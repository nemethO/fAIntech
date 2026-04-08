import os
import json
import hashlib
import time
import pdfplumber
import chardet
import google.generativeai as genai
from openai import OpenAI
from flask import current_app


def extract_text(file_path):
    """Nyers szoveg kinyerese a fajlbol, tipustol fuggetlenul."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.csv':
        return _read_csv_as_text(file_path)
    elif ext == '.pdf':
        return _read_pdf_as_text(file_path)
    elif ext in ('.xlsx', '.xls'):
        return _read_excel_as_text(file_path)
    else:
        raise ValueError(f"Nem tamogatott fajlformatum: {ext}")


def _read_csv_as_text(file_path):
    """CSV fajl beolvasasa, automatikus encoding felismeressel."""
    with open(file_path, 'rb') as f:
        raw = f.read()

    detected = chardet.detect(raw)
    encoding = detected.get('encoding') or 'utf-8'

    # neha a chardet rosszul tippel, probaljuk a gyakori magyar kodolasokat
    for enc in [encoding, 'utf-8', 'iso-8859-2', 'latin-1']:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue

    return raw.decode('utf-8', errors='replace')


def _read_pdf_as_text(file_path):
    """PDF fajl osszes oldalanak szoveget osszefuzi."""
    pages = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n\n".join(pages)


def _read_excel_as_text(file_path):
    """Excel fajl beolvasasa es szovegge alakitasa."""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True)
    lines = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            line = ";".join(str(cell) if cell is not None else "" for cell in row)
            lines.append(line)
    wb.close()
    return "\n".join(lines)


def _call_llm(model, prompt, max_retries=3):
    """LLM hivas ujraprobalkozassal rate limit eseten."""
    for attempt in range(max_retries):
        try:
            return model.generate_content(
                prompt,
                generation_config=genai.GenerationConfig(
                    temperature=0.1,
                    response_mime_type="application/json",
                )
            )
        except Exception as e:
            if '429' in str(e) and attempt < max_retries - 1:
                wait = (attempt + 1) * 5
                time.sleep(wait)
                continue
            raise


def _call_openrouter(prompt, max_retries=3):
    """OpenRouter API hivas, OpenAI-kompatibilis formatum."""
    # 120 masodperces timeout - ingyenes modellek lassuak lehetnek
    client = OpenAI(
        base_url='https://openrouter.ai/api/v1',
        api_key=current_app.config['OPENROUTER_API_KEY'],
        timeout=120.0,
    )
    model_name = current_app.config['OPENROUTER_MODEL']

    for attempt in range(max_retries):
        try:
            print(f"[OpenRouter] Kérés küldése: {model_name} (próba {attempt+1}/{max_retries})...", flush=True)
            import time as _t
            _start = _t.time()
            resp = client.chat.completions.create(
                model=model_name,
                messages=[{'role': 'user', 'content': prompt}],
                temperature=0.1,
                # response_format-ot NEM hasznalunk, mert nem minden modell tamogatja
                # helyette a prompt maga keri a JSON valaszt
            )
            elapsed = _t.time() - _start
            raw = resp.choices[0].message.content
            print(f"[OpenRouter] Válasz érkezett {elapsed:.1f}s alatt ({len(raw)} karakter)", flush=True)

            # neha a modell markdown kod-blokkba teszi a JSON-t, azt le kell szedni
            if raw.startswith('```'):
                raw = raw.split('```')[1]
                if raw.startswith('json'):
                    raw = raw[4:]
            raw = raw.strip()

            class _Resp:
                text = raw
            return _Resp()
        except Exception as e:
            print(f"[OpenRouter] Hiba (próba {attempt+1}): {type(e).__name__}: {e}", flush=True)
            if attempt < max_retries - 1:
                time.sleep((attempt + 1) * 5)
                continue
            raise


def _get_llm_caller():
    """Visszaadja a megfelelo LLM hivorutat a AI_PROVIDER config alapjan."""
    provider = current_app.config.get('AI_PROVIDER', 'gemini')
    if provider == 'openrouter':
        return None, _call_openrouter
    # alapertelmezett: Gemini
    genai.configure(api_key=current_app.config['GEMINI_API_KEY'])
    model = genai.GenerativeModel(current_app.config['GEMINI_MODEL'])
    return model, None


def parse_with_llm(raw_text):
    """
    Nyers banki szoveget kuldunk az LLM-nek, strukturalt tranzakciokat kapunk vissza.
    Parse-olas es kategorizalas EGYETLEN LLM hivasban tortenik a gyorsasag erdekeben.
    Visszaad egy dict-et: { bank_name, currency, transactions: [...] }
    """
    model, openrouter_caller = _get_llm_caller()

    # ha tul hosszu a szoveg, levagjuk (token limit miatt)
    max_chars = 60000
    if len(raw_text) > max_chars:
        raw_text = raw_text[:max_chars]

    prompt = f"""Banki számlakivonatot kapsz nyers szöveg formában. A feladatod:
1. Azonosítsd melyik bank számlakivonata ez
2. Kinyerd az ÖSSZES tranzakciót strukturált formában
3. Minden tranzakcióhoz adj megbízhatósági pontszámot (confidence: 0.0 - 1.0)
4. Minden tranzakciót KATEGORIZÁLJ is az alábbi kategóriák egyikébe

ELÉRHETŐ KATEGÓRIÁK (CSAK ezeket használd):
- Élelmiszer (boltok, szupermarketek: Tesco, Spar, Lidl, Aldi, stb.)
- Étterem (éttermek, gyorsétterem, kávézó: KFC, McDonald's, Pizza King, stb.)
- Közlekedés (benzin, BKK, taxi, Uber, parkolás)
- Szórakozás (mozi, Netflix, Spotify, játékok, koncert)
- Vásárlás (ruha, elektronika, Amazon, webshop, Media Markt)
- Közüzemi díjak (áram, gáz, víz, internet, telefon)
- Lakhatás (albérlet, lakbér, lakáshitel, biztosítás)
- Egészség (gyógyszertár, orvos, kórház)
- Pénzügyi (pénzváltás, átutalás, feltöltés, befektetés)
- Fizetés (munkabér, ösztöndíj, bevétel)
- Egyéb (ami nem illik máshova)

FONTOS SZABÁLYOK:
- A dátumot mindig YYYY-MM-DD formátumra alakítsd
- Az összeg legyen szám (float), negatív ha kiadás, pozitív ha bevétel
- Ha a partner neve nem egyértelmű, írd be amit találsz
- A confidence legyen alacsonyabb ha bizonytalan vagy egy mezőben
- NE találj ki tranzakciókat, csak ami ténylegesen szerepel a szövegben
- Pénzváltásokat, díjakat, feltöltéseket is vedd bele

Válaszolj KIZÁRÓLAG az alábbi JSON formátumban, más szöveget NE írj:
{{
    "bank_name": "Bank neve (pl. OTP, Revolut, Erste, K&H, stb.)",
    "currency": "Fő pénznem (pl. HUF, EUR)",
    "transactions": [
        {{
            "date": "YYYY-MM-DD",
            "amount": -1234.56,
            "currency": "HUF",
            "partner": "Partner/kereskedő neve",
            "description": "Tranzakció leírása/közlemény",
            "is_income": false,
            "confidence": 0.95,
            "category": "Élelmiszer",
            "category_confidence": 0.9
        }}
    ]
}}

A nyers szöveg:
---
{raw_text}
---"""

    response = openrouter_caller(prompt) if openrouter_caller else _call_llm(model, prompt)

    result = json.loads(response.text)

    # hash generalas a duplikaciok elkerulesere
    for tx in result.get('transactions', []):
        hash_input = f"{tx.get('date')}|{tx.get('amount')}|{tx.get('partner', '')}|{tx.get('description', '')}"
        tx['transaction_hash'] = hashlib.sha256(hash_input.encode('utf-8')).hexdigest()
        # ha nincs kategoria, legyen Egyeb
        if 'category' not in tx:
            tx['category'] = 'Egyéb'
            tx['category_confidence'] = 0.5

    return result


def categorize_transactions(transactions):
    """
    Kategorizalas — mar a parse_with_llm megcsinalja egyetlen lepesben,
    ez a fuggveny csak biztonsagi halokent mukodik ha valami kimaradt volna.
    """
    for tx in transactions:
        if 'category' not in tx:
            tx['category'] = 'Egyéb'
            tx['category_confidence'] = 0.5
    return transactions
